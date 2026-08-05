from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


def fail(message: str) -> None:
    raise SystemExit(message)


def main() -> None:
    if len(sys.argv) != 2:
        fail("usage: verify_northwing_office.py <fixture-dir>")
    root = Path(sys.argv[1])

    document = Document(root / "northwing.docx")
    doc_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if "Northwing Report" not in doc_text or "Verified document body." not in doc_text:
        fail(f"DOCX content mismatch: {doc_text!r}")

    presentation = Presentation(root / "northwing.pptx")
    if len(presentation.slides) != 2:
        fail(f"PPTX slide count: {len(presentation.slides)}")
    slide_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    if "First slide" not in slide_text or "Verified bullet" not in slide_text:
        fail(f"PPTX content mismatch: {slide_text!r}")

    workbook = load_workbook(root / "northwing.xlsx", data_only=False)
    if workbook.sheetnames != ["Data"]:
        fail(f"XLSX sheets: {workbook.sheetnames!r}")
    sheet = workbook["Data"]
    if sheet["A2"].value != "Alpha" or sheet["B2"].value != 42:
        fail(f"XLSX values: {sheet['A2'].value!r}, {sheet['B2'].value!r}")

    pdf = PdfReader(root / "northwing.pdf")
    if len(pdf.pages) != 1:
        fail(f"PDF page count: {len(pdf.pages)}")
    pdf_text = pdf.pages[0].extract_text() or ""
    if "Northwing PDF" not in pdf_text or "Verified searchable body." not in pdf_text:
        fail(f"PDF content mismatch: {pdf_text!r}")

    print("Northwing Office fixtures opened successfully with independent parsers")


if __name__ == "__main__":
    main()
